#!/usr/bin/env python3
"""Generate a PowerPoint presentation from a Tungsten corporate .potx template.

Usage:
    echo '{"title":"My Deck","slides":[...]}' | python generate_deck.py path/to/template.potx

Reads JSON from stdin, applies it to the template, and prints the path to the
generated .pptx file on stdout.
"""

import json
import os
import sys
import tempfile

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor

# ---------------------------------------------------------------------------
# Layout-name-to-index mapping (based on template analysis — 23 layouts)
# ---------------------------------------------------------------------------
LAYOUT_MAP = {
    "title": 1,            # Title Slide 3  (idx 0=title, idx 1=subtitle)
    "content-white": 7,    # White - Title and Content  (idx 0=title, idx 14=body)
    "content-blue": 15,    # Blue - Title and Content   (idx 0=title, idx 14=body)
    "section": 20,         # Section Slide  (idx 0=title, idx 14=body)
    "closing": 22,         # Closing Slide  (idx 0=title only)
    # New types use content-white as base layout
    "stats": 7,            # Stats slide — uses white layout, stats rendered in body
    "comparison": 7,       # Comparison table — uses white layout, table added as shape
}

FOOTER_TEXT = (
    "Proprietary and Confidential \u00a9 2026 Tungsten Automation Corp. "
    "All rights reserved."
)

# Brand colors
TUNGSTEN_NAVY = RGBColor(0x00, 0x28, 0x54)
TUNGSTEN_BLUE = RGBColor(0x1E, 0x4D, 0x8C)
TUNGSTEN_GOLD = RGBColor(0xFF, 0xC6, 0x00)
TUNGSTEN_BRIGHT = RGBColor(0x00, 0xA0, 0xFB)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK_TEXT = RGBColor(0x23, 0x1F, 0x20)
GRAY_TEXT = RGBColor(0x6B, 0x72, 0x80)


def _safe_layout(prs, index):
    """Return the slide layout at *index*, falling back to index 0."""
    layouts = prs.slide_layouts
    if 0 <= index < len(layouts):
        return layouts[index]
    return layouts[0]


def _find_body_placeholder(slide):
    """Return the body text placeholder (idx 14) for bullet content."""
    if 14 in slide.placeholders:
        return slide.placeholders[14]
    # Fallback: first non-title, non-slide-number placeholder with a text frame
    for ph in slide.placeholders:
        if ph.placeholder_format.idx not in (0, 1, 10, 11, 12) and ph.has_text_frame:
            return ph
    return None


def _set_title(slide, text):
    """Set the slide title in placeholder 0 or the first text placeholder."""
    if not text:
        return
    if 0 in slide.placeholders:
        slide.placeholders[0].text = text
        return
    for ph in slide.placeholders:
        if ph.has_text_frame:
            ph.text = text
            return


def _set_subtitle(slide, text):
    """Set text in placeholder 1 (subtitle) if it exists."""
    if not text:
        return
    if 1 in slide.placeholders:
        slide.placeholders[1].text = text


def _set_bullets(slide, bullets):
    """Populate the body placeholder with a list of bullet strings."""
    if not bullets:
        return
    body = _find_body_placeholder(slide)
    if body is None:
        return
    tf = body.text_frame
    tf.clear()
    for i, bullet in enumerate(bullets):
        if i == 0:
            tf.paragraphs[0].text = bullet
        else:
            p = tf.add_paragraph()
            p.text = bullet


def _set_stats(slide, stats):
    """Render stats as large bold numbers with descriptors in the body placeholder.

    Pattern inspired by SBR Industry Insights slide:
      >33%  |  90%  |  45%
      desc    desc    desc
    """
    if not stats:
        return
    body = _find_body_placeholder(slide)
    if body is None:
        return
    tf = body.text_frame
    tf.clear()

    for i, stat in enumerate(stats):
        value = stat.get("value", "")
        label = stat.get("label", "")

        # Large bold number
        if i == 0:
            p_val = tf.paragraphs[0]
        else:
            # Add spacing between stats
            p_spacer = tf.add_paragraph()
            p_spacer.space_before = Pt(12)
            p_val = tf.add_paragraph()

        run_val = p_val.add_run()
        run_val.text = value
        run_val.font.size = Pt(36)
        run_val.font.bold = True
        run_val.font.color.rgb = TUNGSTEN_NAVY

        # Smaller descriptor below
        p_label = tf.add_paragraph()
        run_label = p_label.add_run()
        run_label.text = label
        run_label.font.size = Pt(14)
        run_label.font.color.rgb = GRAY_TEXT


def _set_comparison_table(slide, comparison):
    """Render a comparison table as a python-pptx table shape.

    Pattern inspired by Buy vs Build deck:
      Capability  |  Tungsten  |  DIY Build
      Row 1...
    """
    if not comparison:
        return
    headers = comparison.get("headers", [])
    rows = comparison.get("rows", [])
    if not headers or not rows:
        return

    num_cols = len(headers)
    num_rows = len(rows) + 1  # +1 for header row

    # Position table below the title area
    left = Inches(0.8)
    top = Inches(1.8)
    width = Inches(10.5)
    height = Inches(0.4) * num_rows

    table_shape = slide.shapes.add_table(num_rows, num_cols, left, top, width, height)
    table = table_shape.table

    # Style header row
    for col_idx, header_text in enumerate(headers):
        cell = table.cell(0, col_idx)
        cell.text = header_text
        for paragraph in cell.text_frame.paragraphs:
            for run in paragraph.runs:
                run.font.bold = True
                run.font.size = Pt(12)
                run.font.color.rgb = WHITE
            paragraph.alignment = PP_ALIGN.LEFT
        # Navy background for header
        cell.fill.solid()
        cell.fill.fore_color.rgb = TUNGSTEN_NAVY

    # Data rows — alternate white and light gray
    for row_idx, row_data in enumerate(rows):
        for col_idx, cell_text in enumerate(row_data):
            if col_idx >= num_cols:
                break
            cell = table.cell(row_idx + 1, col_idx)
            cell.text = str(cell_text)
            for paragraph in cell.text_frame.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(11)
                    run.font.color.rgb = DARK_TEXT
                paragraph.alignment = PP_ALIGN.LEFT
            # Alternate row shading
            if row_idx % 2 == 1:
                cell.fill.solid()
                cell.fill.fore_color.rgb = RGBColor(0xF0, 0xF0, 0xF0)

    # Set column widths — first column wider
    if num_cols >= 2:
        col_width_first = int(width.emu * 0.4)
        col_width_rest = int(width.emu * 0.6 / (num_cols - 1))
        table.columns[0].width = col_width_first
        for i in range(1, num_cols):
            table.columns[i].width = col_width_rest


def _set_speaker_notes(slide, notes_text):
    """Add speaker notes to a slide."""
    if not notes_text:
        return
    notes_slide = slide.notes_slide
    notes_slide.notes_text_frame.text = notes_text


def _set_footer(slide, text):
    """Set the footer placeholder text on a slide, if the placeholder exists."""
    for ph in slide.placeholders:
        try:
            if ph.placeholder_format.idx in (11, 12):
                ph.text = text
                return
        except Exception:
            continue


def _remove_all_slides(prs):
    """Remove all pre-existing slides from the presentation.

    Handles multiple python-pptx versions safely.
    """
    while len(prs.slides) > 0:
        rId = None
        slide = prs.slides[0]
        for key, rel in prs.part.rels.items():
            try:
                if rel.target_part is slide.part:
                    rId = key
                    break
            except Exception:
                continue
        if rId:
            prs.part.drop_rel(rId)

        # Remove from sldIdLst — try multiple approaches
        try:
            sldIdLst = prs.part._element.sldIdLst
            if sldIdLst is not None and len(sldIdLst) > 0:
                sldIdLst.remove(sldIdLst[0])
                continue
        except AttributeError:
            pass

        # Fallback: use lxml to find and remove sldId elements
        try:
            from lxml import etree
            nsmap = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main"}
            sldIdLst_el = prs.part._element.find(".//p:sldIdLst", nsmap)
            if sldIdLst_el is not None and len(sldIdLst_el) > 0:
                sldIdLst_el.remove(sldIdLst_el[0])
                continue
        except Exception:
            pass

        # Last resort: break to avoid infinite loop
        print("Warning: could not remove all template slides", file=sys.stderr)
        break


def _open_template(template_path):
    """Open a .potx or .pptx template, patching content-type if needed."""
    import zipfile, shutil

    is_potx = template_path.lower().endswith(".potx")

    if is_potx:
        # python-pptx rejects .potx — rewrite [Content_Types].xml
        fd_tpl, tmp_tpl = tempfile.mkstemp(suffix=".pptx")
        os.close(fd_tpl)
        with zipfile.ZipFile(template_path, "r") as src_zip:
            with zipfile.ZipFile(tmp_tpl, "w", zipfile.ZIP_DEFLATED) as dst_zip:
                for item in src_zip.infolist():
                    raw = src_zip.read(item.filename)
                    if item.filename == "[Content_Types].xml":
                        raw = raw.replace(
                            b"application/vnd.openxmlformats-officedocument.presentationml.template.main+xml",
                            b"application/vnd.openxmlformats-officedocument.presentationml.presentation.main+xml",
                        )
                    dst_zip.writestr(item, raw)
        prs = Presentation(tmp_tpl)
        os.unlink(tmp_tpl)
    else:
        # .pptx can be opened directly, but we copy to temp first to avoid
        # modifying the original
        fd_tpl, tmp_tpl = tempfile.mkstemp(suffix=".pptx")
        os.close(fd_tpl)
        shutil.copy2(template_path, tmp_tpl)
        prs = Presentation(tmp_tpl)
        os.unlink(tmp_tpl)

    return prs


def _build_slide(prs, slide_def):
    """Add a single slide to the presentation based on slide_def JSON."""
    layout_name = slide_def.get("layout", "content-white")
    layout_idx = LAYOUT_MAP.get(layout_name, LAYOUT_MAP["content-white"])
    layout = _safe_layout(prs, layout_idx)

    slide = prs.slides.add_slide(layout)

    _set_title(slide, slide_def.get("title"))

    subtitle = slide_def.get("subtitle")
    if layout_name == "title":
        _set_subtitle(slide, subtitle)
    elif layout_name in ("section", "closing") and subtitle:
        body = _find_body_placeholder(slide)
        if body is not None:
            body.text = subtitle

    if layout_name == "stats":
        _set_stats(slide, slide_def.get("stats"))
    elif layout_name == "comparison":
        _set_comparison_table(slide, slide_def.get("comparison"))
    elif layout_name in ("content-white", "content-blue"):
        _set_bullets(slide, slide_def.get("bullets"))

    _set_speaker_notes(slide, slide_def.get("speakerNotes"))
    return slide


def _reorder_slides(prs, new_order):
    """Reorder slides in the presentation to match new_order (list of 0-based indices).

    Uses low-level XML manipulation on the sldIdLst element.
    """
    try:
        sldIdLst = prs.part._element.sldIdLst
        if sldIdLst is None:
            return
        sldId_elements = list(sldIdLst)
        # Clear and re-add in new order
        for el in sldId_elements:
            sldIdLst.remove(el)
        for idx in new_order:
            if idx < len(sldId_elements):
                sldIdLst.append(sldId_elements[idx])
    except Exception as exc:
        print(f"Warning: slide reorder failed: {exc}", file=sys.stderr)


def main():
    # ---- Parse arguments ----------------------------------------------------
    if len(sys.argv) < 2:
        print("Usage: echo '{...}' | python generate_deck.py <template> [--mode <mode>]",
              file=sys.stderr)
        sys.exit(1)

    template_path = sys.argv[1]

    # Parse --mode flag
    mode = "replace-all"  # default: current behavior
    if "--mode" in sys.argv:
        mode_idx = sys.argv.index("--mode")
        if mode_idx + 1 < len(sys.argv):
            mode = sys.argv[mode_idx + 1]

    # ---- Read JSON from stdin -----------------------------------------------
    try:
        import io
        stdin_text = io.TextIOWrapper(sys.stdin.buffer, encoding="utf-8").read()
        data = json.loads(stdin_text)
    except (json.JSONDecodeError, ValueError) as exc:
        print(f"Invalid JSON on stdin: {exc}", file=sys.stderr)
        sys.exit(1)

    # ---- Open template ------------------------------------------------------
    try:
        prs = _open_template(template_path)
    except Exception as exc:
        print(f"Failed to open template: {exc}", file=sys.stderr)
        sys.exit(1)

    slides_data = data.get("slides", [])

    # ---- Mode: replace-all (Slides Only — current behavior) -----------------
    if mode == "replace-all":
        _remove_all_slides(prs)
        for slide_def in slides_data:
            _build_slide(prs, slide_def)

    # ---- Mode: insert-after-N (L0 or SBR — insert into existing deck) ------
    elif mode.startswith("insert-after-"):
        insert_after = int(mode.split("-")[-1])  # e.g., 1 or 3
        existing_count = len(prs.slides)

        # Add new slides at the end first (python-pptx only appends)
        for slide_def in slides_data:
            _build_slide(prs, slide_def)

        # Now reorder: [0..insert_after-1] + [new slides] + [insert_after..existing-1]
        total = len(prs.slides)
        num_new = len(slides_data)

        # Original slide indices: 0 to existing_count-1
        # New slide indices: existing_count to total-1
        before = list(range(0, insert_after))                 # slides before insertion point
        new = list(range(existing_count, total))              # newly added slides
        after = list(range(insert_after, existing_count))     # remaining original slides

        new_order = before + new + after
        _reorder_slides(prs, new_order)

    else:
        print(f"Unknown mode: {mode}", file=sys.stderr)
        sys.exit(1)

    # ---- Set footer on every slide ------------------------------------------
    for slide in prs.slides:
        _set_footer(slide, FOOTER_TEXT)

    # ---- Save to temp file --------------------------------------------------
    fd, tmp_path = tempfile.mkstemp(suffix=".pptx")
    os.close(fd)
    try:
        prs.save(tmp_path)
    except Exception as exc:
        print(f"Failed to save presentation: {exc}", file=sys.stderr)
        sys.exit(1)

    print(tmp_path)


if __name__ == "__main__":
    main()
