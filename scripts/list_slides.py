from pptx import Presentation
import os

templates_dir = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "templates")

for name in ['TungstenL0.pptx', 'TungstenSBR.pptx']:
    path = os.path.join(templates_dir, name)
    prs = Presentation(path)
    print(f"\n{name}: {len(prs.slides)} slides")
    for i, slide in enumerate(prs.slides):
        texts = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                t = shape.text_frame.text.strip()[:80]
                if t and "Proprietary" not in t:
                    texts.append(t)
        summary = " | ".join(texts[:3]) if texts else "(images/graphics only)"
        print(f"  Slide {i+1}: {summary}")
